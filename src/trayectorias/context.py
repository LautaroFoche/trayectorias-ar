"""Adaptadores de contexto: WDI, PISA y extracción editorial verificable de Kids Online."""

from __future__ import annotations
from concurrent.futures import ThreadPoolExecutor
from datetime import datetime, timezone
import io
import json
import math
from pathlib import Path
import subprocess
import tempfile

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from .ingest import ContractError, atomic_json, sha256

COUNTRIES = {
    "ARG": "Argentina",
    "BRA": "Brasil",
    "CHL": "Chile",
    "URY": "Uruguay",
    "OECD": "Promedio OECD",
}
PISA_COUNTRIES = {
    "Argentina": "ARG",
    "Brazil": "BRA",
    "Chile": "CHL",
    "Uruguay": "URY",
    "OECD average": "OECD",
}
CONTEXT_COLUMNS = {
    "observation_id": "VARCHAR",
    "source_id": "VARCHAR",
    "country_code": "VARCHAR",
    "country_name": "VARCHAR",
    "year": "INTEGER",
    "publication_year": "INTEGER",
    "indicator": "VARCHAR",
    "indicator_label": "VARCHAR",
    "unit": "VARCHAR",
    "population": "VARCHAR",
    "subgroup": "VARCHAR",
    "category": "VARCHAR",
    "category_order": "INTEGER",
    "value": "DOUBLE",
    "standard_error": "DOUBLE",
    "value_status": "VARCHAR",
    "source_flag": "VARCHAR",
    "evidence_kind": "VARCHAR",
    "denominator": "VARCHAR",
    "sample_n": "INTEGER",
    "method_note": "VARCHAR",
    "source_locator": "VARCHAR",
    "source_sha256": "VARCHAR",
    "source_url": "VARCHAR",
    "extraction_method": "VARCHAR",
}
SOURCE_COLUMNS = {
    k: "VARCHAR"
    for k in [
        "source_id",
        "family",
        "producer",
        "url",
        "sha256",
        "path",
        "retrieved_at",
        "format",
    ]
}
SOURCE_COLUMNS["bytes"] = "BIGINT"


def fetch_context(root, offline=False):
    specs = json.loads((root / "config/context_sources.json").read_text())
    path = root / "data/raw/context_manifest.json"
    if offline:
        manifest = json.loads(path.read_text())
        if {(s["id"], s["url"]) for s in specs} != {
            (s["id"], s["url"]) for s in manifest
        }:
            raise ContractError("Contexto offline no coincide con la configuración")
        configured = {s["id"]: s for s in specs}
        for source in manifest:
            if sha256((root / source["path"]).read_bytes()) != source["sha256"]:
                raise ContractError("Hash inválido: " + source["id"])
            if (
                configured[source["id"]].get("reviewed_sha256", source["sha256"])
                != source["sha256"]
            ):
                raise ContractError("PDF pendiente de revisión: " + source["id"])
        return manifest

    def fetch(spec):
        with tempfile.TemporaryDirectory() as folder:
            dest = Path(folder) / "download"
            result = subprocess.run(
                [
                    "curl",
                    "--fail",
                    "--silent",
                    "--show-error",
                    "--location",
                    "--proto",
                    "=https",
                    "--proto-redir",
                    "=https",
                    "--retry",
                    "2",
                    "--retry-delay",
                    "1",
                    "--max-time",
                    "90",
                    "--max-filesize",
                    "50000000",
                    spec["url"],
                    "-o",
                    str(dest),
                ],
                capture_output=True,
                text=True,
            )
            if result.returncode:
                raise ContractError(
                    "Falló descarga " + spec["id"] + ": " + result.stderr.strip()
                )
            content = dest.read_bytes()
        digest = sha256(content)
        if spec.get("reviewed_sha256", digest) != digest:
            raise ContractError(
                "El PDF cambió; revisar páginas/denominadores antes de aceptar "
                + spec["id"]
            )
        if spec["format"] == "json":
            json.loads(content)  # rechazar páginas HTML de error
        if spec["format"] == "pdf" and not content.startswith(b"%PDF"):
            raise ContractError("PDF inválido")
        if spec["format"] == "xlsx" and not content.startswith(b"PK"):
            raise ContractError("XLSX inválido")
        relative = (
            Path("data/raw/context") / spec["id"] / (digest + "." + spec["format"])
        )
        saved = root / relative
        saved.parent.mkdir(parents=True, exist_ok=True)
        if saved.exists() and sha256(saved.read_bytes()) != digest:
            raise ContractError("Snapshot alterado")
        if not saved.exists():
            saved.write_bytes(content)
        return dict(
            spec,
            path=str(relative),
            sha256=digest,
            bytes=len(content),
            retrieved_at=datetime.now(timezone.utc).isoformat(),
        )

    # Sólo descargas independientes. Publicación del manifiesto después de completar todas.
    with ThreadPoolExecutor(max_workers=4) as pool:
        manifest = list(pool.map(fetch, specs))
    atomic_json(path, manifest)
    return manifest


def observation(source, **fields):
    row = dict.fromkeys(CONTEXT_COLUMNS)
    row.update(
        source_id=source["id"],
        source_sha256=source["sha256"],
        source_url=source["url"],
        subgroup="Total",
        category="Total",
        category_order=0,
        standard_error=None,
        source_flag="",
        sample_n=None,
        extraction_method="automatico",
        publication_year=None,
        method_note="",
    )
    row.update(fields)
    value = row["value"]
    if value is not None and (
        isinstance(value, bool)
        or not isinstance(value, (int, float))
        or not math.isfinite(value)
    ):
        raise ContractError("Valor de contexto inválido")
    se = row["standard_error"]
    if se is not None and (
        not isinstance(se, (int, float)) or not math.isfinite(se) or se < 0
    ):
        raise ContractError("Error estándar inválido")
    row["country_name"] = COUNTRIES[row["country_code"]]
    row["value_status"] = "sin_dato" if value is None else "observado"
    key = [
        row[k]
        for k in [
            "source_id",
            "country_code",
            "year",
            "indicator",
            "subgroup",
            "category",
        ]
    ]
    row["observation_id"] = sha256(json.dumps(key, ensure_ascii=False).encode())
    return row


def parse_wdi(content, source, specs):
    body = json.loads(content)
    if not isinstance(body, list) or len(body) != 2 or not isinstance(body[0], dict):
        raise ContractError("Respuesta WDI inválida")
    meta, rows = body
    if (
        int(meta.get("pages", 0)) != 1
        or int(meta.get("page", 0)) != 1
        or meta.get("total") != len(rows)
    ):
        raise ContractError(
            "Paginación WDI incompleta: no publicar una respuesta parcial"
        )
    expected = {
        (c, m["code"], y)
        for c in COUNTRIES
        if c != "OECD"
        for m in specs
        for y in range(2000, 2026)
    }
    lookup = {m["code"]: m for m in specs}
    seen = set()
    records = []
    for i, item in enumerate(rows):
        country = item["countryiso3code"]
        code = item["indicator"]["id"]
        year = int(item["date"])
        key = (country, code, year)
        if key not in expected or key in seen:
            raise ContractError("Cobertura/duplicación WDI inesperada")
        seen.add(key)
        spec = lookup[code]
        value = item["value"]
        if value is not None and value < 0:
            raise ContractError("WDI negativo inesperado")
        if code == "IT.NET.USER.ZS" and value is not None and value > 100:
            raise ContractError("Porcentaje de personas inválido")
        records.append(
            observation(
                source,
                country_code=country,
                year=year,
                indicator=code,
                indicator_label=spec["label"],
                unit=spec["unit"],
                value=value,
                population="Agregado nacional; población según definición del indicador",
                evidence_kind="serie_agregada",
                denominator=spec["unit"],
                method_note=spec["note"],
                source_locator=f"JSON /1/{i}/value; {country}/{code}/{year}",
                source_flag=item.get("obs_status", ""),
            )
        )
    if seen != expected:
        raise ContractError("WDI no cubre el contrato país × indicador × año")
    return records


def pisa_number(value):
    if (
        isinstance(value, (float, int))
        and not isinstance(value, bool)
        and math.isfinite(value)
    ):
        return float(value), ""
    if value in ("m", "c", "a", "w", "x", None, ""):
        return None, str(value or "vacío")
    raise ContractError("Marca PISA desconocida: " + repr(value))


def parse_pisa(content, source):
    book = load_workbook(io.BytesIO(content), data_only=True, read_only=True)
    records = []
    sheets = (
        ["Table II.B1.3.9"]
        if source["id"] == "pisa_ch3"
        else ["Table II.B1.5.64", "Table II.B1.5.65", "Table II.B1.5.67"]
    )
    categories = [
        "Sin uso",
        "Hasta 1 h",
        "Más de 1–2 h",
        "Más de 2–3 h",
        "Más de 3–5 h",
        "Más de 5–7 h",
        "Más de 7 h",
    ]
    original_categories = [
        "None",
        "Up to 1 hour",
        "More than 1 hour and up to 2 hours",
        "More than 2 hours and up to 3 hours",
        "More than 3 hours and up to 5 hours",
        "More than 5 hours and up to 7 hours",
        "More than 7 hours",
    ]
    for name in sheets:
        if name not in book.sheetnames:
            raise ContractError("Falta tabla PISA " + name)
        rows = list(book[name].values)

        def cell(r, c):
            return rows[r - 1][c - 1]

        if cell(1, 1) != name:
            raise ContractError("Tabla PISA no reconocida")
        selections = []
        if name.endswith(".3.9"):
            for col, indicator, label, header in [
                (
                    68,
                    "distraccion_propia",
                    "Distracción por el propio dispositivo",
                    "Students get distracted by using digital devices",
                ),
                (
                    80,
                    "distraccion_ajena",
                    "Distracción por dispositivos de otros",
                    "Students get distracted by other students who are using digital devices",
                ),
            ]:
                if cell(7, col) != header:
                    raise ContractError("Cambió pregunta PISA")
                for offset, (original, category) in enumerate(
                    zip(
                        [
                            "Every lesson",
                            "Most lessons",
                            "Some lessons",
                            "Never or hardly ever",
                        ],
                        [
                            "Todas las clases",
                            "Mayoría de las clases",
                            "Algunas clases",
                            "Nunca o casi nunca",
                        ],
                    )
                ):
                    c = col + offset * 3
                    if (
                        cell(8, c) != original
                        or cell(9, c) != "%"
                        or cell(9, c + 1) != "S.E."
                    ):
                        raise ContractError("Cambió encabezado PISA")
                    selections.append(
                        (
                            c,
                            indicator,
                            label,
                            category,
                            offset,
                            "%",
                            "autorreporte_descriptivo",
                            "Frecuencia en clases de matemática",
                        )
                    )
        elif name.endswith((".5.64", ".5.65")):
            learning = name.endswith(".5.64")
            purpose = "aprendizaje" if learning else "ocio"
            if ("learning" if learning else "leisure") not in str(cell(6, 2)):
                raise ContractError("Cambió propósito de uso PISA")
            for i, (original, category) in enumerate(
                zip(original_categories, categories)
            ):
                c = 2 + i * 3
                if (
                    cell(8, c) != original
                    or cell(9, c) != "Mean score"
                    or cell(9, c + 1) != "S.E."
                ):
                    raise ContractError("Cambió escala PISA")
                selections.append(
                    (
                        c,
                        "matematica_" + purpose,
                        "Matemática según horas de "
                        + purpose
                        + " digital en la escuela",
                        category,
                        i,
                        "puntos PISA",
                        "asociacion_no_ajustada",
                        "Medias publicadas, sin ajuste socioeconómico; no efecto causal",
                    )
                )
        else:
            if "one-hour increase" not in str(cell(6, 2)):
                raise ContractError("Cambió regresión PISA")
            for c, category, index, anchor in [
                (2, "Sin ajuste", 0, "Before accounting"),
                (5, "Ajuste socioeconómico", 1, "After accounting"),
            ]:
                if (
                    anchor not in str(cell(8, c))
                    or cell(9, c) != "Score dif."
                    or cell(9, c + 1) != "S.E."
                ):
                    raise ContractError("Cambió modelo PISA")
                selections.append(
                    (
                        c,
                        "asociacion_ocio",
                        "Asociación entre una hora adicional de ocio digital escolar y matemática",
                        category,
                        index,
                        "puntos PISA / hora",
                        "asociacion_observacional",
                        "Ajuste ESCS de estudiantes y escuelas en categoría ajustada; asociación, no efecto causal",
                    )
                )
        found = set()
        for number, row in enumerate(rows, 1):
            if row[0] not in PISA_COUNTRIES:
                continue
            country = PISA_COUNTRIES[row[0]]
            if country in found:
                raise ContractError("País PISA duplicado")
            found.add(country)
            for col, indicator, label, category, order, unit, kind, note in selections:
                value, flag = pisa_number(row[col - 1])
                se, seflag = pisa_number(row[col])
                if unit == "%" and value is not None and not 0 <= value <= 100:
                    raise ContractError("PISA porcentaje inválido")
                if value is not None and se is None:
                    raise ContractError("PISA sin error estándar")
                records.append(
                    observation(
                        source,
                        country_code=country,
                        year=2022,
                        publication_year=2023,
                        indicator=indicator,
                        indicator_label=label,
                        unit=unit,
                        value=value,
                        standard_error=se,
                        population="Estudiantes de 15 años escolarizados, universo PISA 2022",
                        evidence_kind=kind,
                        denominator="Estudiantes PISA con respuesta válida; estimaciones ponderadas",
                        category=category,
                        category_order=order,
                        method_note=note,
                        source_locator=f"{name}!{get_column_letter(col)}{number}; SE {get_column_letter(col + 1)}{number}",
                        source_flag=flag,
                    )
                )
        if found != set(COUNTRIES):
            raise ContractError("Falta país o promedio OECD en PISA")
    book.close()
    return records


def parse_kids(root, manifest, registry):
    sources = {s["id"]: s for s in manifest}
    pages = {}
    records = []
    for item in registry["observations"]:
        source = sources[item["source_id"]]
        if source["sha256"] != source["reviewed_sha256"]:
            raise ContractError("PDF sin revisión")
        if source["id"] not in pages:
            text = subprocess.run(
                ["pdftotext", "-layout", str(root / source["path"]), "-"],
                check=True,
                capture_output=True,
                text=True,
            ).stdout
            pages[source["id"]] = text.split("\f")
        page = pages[source["id"]][item["page"] - 1]
        if item["anchor"] not in " ".join(page.split()):
            raise ContractError("Ancla PDF no encontrada: " + item["indicator"])
        if not 0 <= item["value"] <= 100:
            raise ContractError("Porcentaje Kids inválido")
        records.append(
            observation(
                source,
                country_code="ARG",
                year=registry["fieldwork_year"],
                publication_year=registry["publication_year"],
                indicator=item["indicator"],
                indicator_label=item["indicator_label"],
                unit="%",
                value=item["value"],
                subgroup=item["subgroup"],
                category=item["category"],
                evidence_kind=item["evidence_kind"],
                denominator=item["denominator"],
                sample_n=item["sample_n"],
                population=registry["population"],
                method_note=registry["methodology"]
                + " Porcentajes redondeados publicados; sin error estándar disponible en esta extracción.",
                source_locator=f"PDF página {item['page']}; indicador {item['indicator']}; subgrupo {item['subgroup']}",
                extraction_method="revision_manual_con_hash_y_ancla",
            )
        )
    return records


def normalize_context(root, manifest):
    specs = json.loads((root / "config/wdi_indicators.json").read_text())
    records = []
    metadata = []
    for source in manifest:
        content = (root / source["path"]).read_bytes()
        if sha256(content) != source["sha256"]:
            raise ContractError("Contexto alterado antes de cargar")
        if source["kind"] == "wdi":
            records += parse_wdi(content, source, specs)
        elif source["kind"] == "pisa":
            records += parse_pisa(content, source)
        elif source["kind"] == "wdi_metadata":
            data = json.loads(content)
            if (
                not isinstance(data, list)
                or len(data) != 2
                or len(data[1]) != 1
                or data[1][0]["id"] != source["indicator"]
            ):
                raise ContractError("Metadatos WDI incompletos")
            metadata += data[1]
    registry = json.loads((root / "config/kids_reviewed.json").read_text())
    records += parse_kids(root, manifest, registry)
    ids = [r["observation_id"] for r in records]
    if len(ids) != len(set(ids)):
        raise ContractError("Contexto duplicado")
    records.sort(key=lambda r: r["observation_id"])
    atomic_json(root / "data/normalized/context.json", records)
    atomic_json(root / "data/normalized/wdi_metadata.json", metadata)
    return records


def source_catalog(education_manifest, context_manifest):
    rows = []
    for source in education_manifest:
        rows.append(
            dict(
                source_id="educacion_" + source["indicator"],
                family="Secretaría de Educación",
                producer="Secretaría de Educación · RA",
                format="zip",
                **{
                    k: source[k]
                    for k in ["url", "sha256", "path", "retrieved_at", "bytes"]
                },
            )
        )
    for source in context_manifest:
        rows.append(
            dict(
                source_id=source["id"],
                **{
                    k: source[k]
                    for k in [
                        "family",
                        "producer",
                        "format",
                        "url",
                        "sha256",
                        "path",
                        "retrieved_at",
                        "bytes",
                    ]
                },
            )
        )
    return rows
