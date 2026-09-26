"""Snapshots verificables, descarga con reintentos y contrato de Excel cerrado."""

from __future__ import annotations
import hashlib
import io
import json
import math
import os
from pathlib import Path
import re
import subprocess
import tempfile
import unicodedata
import zipfile
from datetime import datetime, timezone

import openpyxl
from openpyxl.utils import get_column_letter

BASE_PAGE = "https://www.argentina.gob.ar/educacion/evaluacion-e-informacion-educativa/indicadores"
GEOGRAPHIES = [
    ("00", "Total País", "pais", None),
    ("02", "Ciudad de Buenos Aires", "provincia", None),
    ("06", "Buenos Aires", "provincia", None),
    ("10", "Catamarca", "provincia", None),
    ("14", "Córdoba", "provincia", None),
    ("18", "Corrientes", "provincia", None),
    ("22", "Chaco", "provincia", None),
    ("26", "Chubut", "provincia", None),
    ("30", "Entre Ríos", "provincia", None),
    ("34", "Formosa", "provincia", None),
    ("38", "Jujuy", "provincia", None),
    ("42", "La Pampa", "provincia", None),
    ("46", "La Rioja", "provincia", None),
    ("50", "Mendoza", "provincia", None),
    ("54", "Misiones", "provincia", None),
    ("58", "Neuquén", "provincia", None),
    ("62", "Río Negro", "provincia", None),
    ("66", "Salta", "provincia", None),
    ("70", "San Juan", "provincia", None),
    ("74", "San Luis", "provincia", None),
    ("78", "Santa Cruz", "provincia", None),
    ("82", "Santa Fe", "provincia", None),
    ("86", "Santiago del Estero", "provincia", None),
    ("90", "Tucumán", "provincia", None),
    ("94", "Tierra del Fuego", "provincia", None),
    ("06C", "Conurbano", "subprovincia", "06"),
    ("06R", "Resto de Bs As", "subprovincia", "06"),
]


class ContractError(ValueError):
    """Un cambio de fuente requiere revisión, nunca corrección silenciosa."""


def normalize_label(value):
    text = unicodedata.normalize("NFKD", str(value))
    return " ".join(
        "".join(c for c in text if not unicodedata.combining(c)).casefold().split()
    )


GEO = {
    normalize_label(name): (code, name, kind, parent)
    for code, name, kind, parent in GEOGRAPHIES
}


def sha256(data):
    return hashlib.sha256(data).hexdigest()


def atomic_json(path: Path, value):
    path.parent.mkdir(parents=True, exist_ok=True)
    tmp = path.with_suffix(path.suffix + ".tmp")
    tmp.write_text(
        json.dumps(value, ensure_ascii=False, indent=2) + "\n", encoding="utf-8"
    )
    os.replace(tmp, path)


def fetch_sources(root: Path, offline=False):
    """Actualiza el manifiesto sólo después de descargar/verificar todas las fuentes."""
    path = root / "data/raw/manifest.json"
    sources = json.loads((root / "config/sources.json").read_text())
    if offline:
        manifest = json.loads(path.read_text())
        if {(s["indicator"], s["url"]) for s in sources} != {
            (s["indicator"], s["url"]) for s in manifest
        }:
            raise ContractError("Configuración y manifiesto offline no coinciden")
        for source in manifest:
            if sha256((root / source["path"]).read_bytes()) != source["sha256"]:
                raise ContractError(f"Hash inválido: {source['indicator']}")
        return manifest
    manifest = []
    for source in sources:
        with tempfile.TemporaryDirectory() as folder:
            dest = Path(folder) / "source.zip"
            # curl funciona con este servidor; urllib sin su User-Agent recibe HTTP 403.
            subprocess.run(
                [
                    "curl",
                    "--fail",
                    "--silent",
                    "--show-error",
                    "--location",
                    "--proto",
                    "=https",
                    "--proto-redir",
                    "=https",
                    "--retry",
                    "3",
                    "--retry-delay",
                    "2",
                    "--max-time",
                    "90",
                    "--max-filesize",
                    "50000000",
                    source["url"],
                    "-o",
                    str(dest),
                ],
                check=True,
            )
            content = dest.read_bytes()
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            if archive.testzip() is not None:
                raise ContractError("ZIP corrupto")
        digest = sha256(content)
        relative = Path("data/raw") / source["indicator"] / (digest + ".zip")
        saved = root / relative
        saved.parent.mkdir(parents=True, exist_ok=True)
        if saved.exists() and sha256(saved.read_bytes()) != digest:
            raise ContractError("Snapshot existente alterado")
        if not saved.exists():
            saved.write_bytes(content)
        manifest.append(
            dict(
                source,
                path=str(relative),
                sha256=digest,
                bytes=len(content),
                retrieved_at=datetime.now(timezone.utc).isoformat(),
            )
        )
    atomic_json(path, manifest)
    return manifest


def parse_workbook(content: bytes, source: dict):
    """Grano: año base × jurisdicción × nivel × grado × indicador."""
    archive = zipfile.ZipFile(io.BytesIO(content))
    candidates = [
        n for n in archive.namelist() if n.endswith(".xlsx") and "territorial" in n
    ]
    if len(candidates) != 1:
        raise ContractError("Se esperaba un libro con estructura provincial")
    member = candidates[0]
    if archive.getinfo(member).file_size > 50_000_000:
        raise ContractError("Libro demasiado grande")
    workbook = openpyxl.load_workbook(io.BytesIO(archive.read(member)), data_only=True)
    records = []
    notes = []
    years = []
    for sheet in workbook:
        if sheet.title == "Limitaciones":
            notes.append(
                {
                    "indicator": source["indicator"],
                    "sheet": sheet.title,
                    "text": " ".join(
                        str(v) for row in sheet.values for v in row if v is not None
                    ),
                }
            )
            continue
        if not re.fullmatch(r"20\d{2}(?:-20\d{2})?", sheet.title):
            raise ContractError(f"Hoja inesperada: {sheet.title}")
        year = int(sheet.title[:4])
        years.append(year)
        if source["indicator"] == "abandono" and sheet.title != f"{year}-{year + 1}":
            raise ContractError("Abandono debe indicar años consecutivos")
        rows = list(sheet.values)
        headers = [
            i for i, r in enumerate(rows) if "Primaria" in r and "Secundaria" in r
        ]
        if len(headers) != 1:
            raise ContractError(f"Encabezado ambiguo en {sheet.title}")
        h = headers[0]
        if rows[h][2] != "Primaria" or rows[h][10] != "Secundaria":
            raise ContractError("Cambió la distribución de columnas")
        expected = (
            ["Total"]
            + [f"{i}° Año" for i in range(1, 8)]
            + ["Total"]
            + [f"{i}° Año" for i in range(7, 13)]
        )
        if list(rows[h + 1][2:17]) != expected:
            raise ContractError("Cambió el encabezado de grados")
        seen = set()
        in_table = False
        for index, row in enumerate(rows[h + 2 :], h + 3):
            label = normalize_label(row[0])
            is_data = row[1] in ("6-6", "7-5") or label in GEO
            if not is_data:
                if row[0] is not None:
                    if not in_table or (isinstance(row[2], (int, float))):
                        raise ContractError(f"Fila desconocida: {row[0]}")
                    notes.append(
                        {
                            "indicator": source["indicator"],
                            "sheet": sheet.title,
                            "text": str(row[0]),
                        }
                    )
                continue
            if label not in GEO:
                raise ContractError(f"Jurisdicción desconocida: {row[0]}")
            code, name, kind, parent = GEO[label]
            if code in seen:
                raise ContractError(f"Jurisdicción duplicada: {code}")
            seen.add(code)
            in_table = True
            structure = row[1] or "mixta"
            if structure not in ("6-6", "7-5", "mixta"):
                raise ContractError("Estructura educativa desconocida")
            for col in range(2, 17):
                level = "primaria" if col < 10 else "secundaria"
                grade = 0 if col in (2, 10) else (col - 2 if col < 10 else col - 4)
                value = row[col]
                structural = (col == 9 and structure == "6-6") or (
                    col == 11 and structure == "7-5"
                )
                if structural and value is not None:
                    raise ContractError("Valor en grado no aplicable")
                if value is not None and (
                    isinstance(value, bool)
                    or not isinstance(value, (int, float))
                    or not math.isfinite(value)
                ):
                    raise ContractError(
                        f"Valor no numérico {sheet.title}:{index}:{col + 1}: {value!r}"
                    )
                status = (
                    "no_aplica"
                    if structural
                    else (
                        "faltante"
                        if value is None
                        else (
                            "fuera_rango_teorico"
                            if value < 0 or value > 100
                            else "observado"
                        )
                    )
                )
                key = f"{year}|{code}|{level}|{grade}|{source['indicator']}"
                records.append(
                    dict(
                        observation_id=sha256(key.encode()),
                        year=year,
                        year_to=year + 1,
                        geo_code=code,
                        geo_name=name,
                        geo_type=kind,
                        parent_code=parent,
                        structure=structure,
                        level=level,
                        grade=grade,
                        indicator=source["indicator"],
                        value=value,
                        value_status=status,
                        source_sha256=source["sha256"],
                        source_url=source["url"],
                        source_member=member,
                        source_sheet=sheet.title,
                        source_cell=f"{get_column_letter(col + 1)}{index}",
                    )
                )
        if seen != set(g[0] for g in GEOGRAPHIES):
            raise ContractError(f"Cobertura incompleta en {sheet.title}: {seen}")
    if not years or len(set(years)) != len(years):
        raise ContractError("Años vacíos/duplicados")
    if sorted(years) != list(range(min(years), max(years) + 1)):
        raise ContractError("Serie con huecos anuales")
    return records, notes


def normalize_sources(root: Path, manifest):
    records = []
    notes = []
    coverage = []
    for source in manifest:
        content = (root / source["path"]).read_bytes()
        if sha256(content) != source["sha256"]:
            raise ContractError("Snapshot alterado antes de normalizar")
        rows, source_notes = parse_workbook(content, source)
        records += rows
        notes += source_notes
        coverage.append(set(r["year"] for r in rows))
    if any(years != coverage[0] for years in coverage):
        raise ContractError("Períodos incompatibles entre indicadores")
    keys = [r["observation_id"] for r in records]
    if len(keys) != len(set(keys)):
        raise ContractError("Claves duplicadas")
    records.sort(key=lambda r: r["observation_id"])
    atomic_json(root / "data/normalized/observations.json", records)
    atomic_json(root / "data/normalized/source_notes.json", notes)
    return records
